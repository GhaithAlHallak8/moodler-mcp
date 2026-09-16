from typing import Any

_TEXT_SUFFIXES = {
    ".py",
    ".txt",
    ".csv",
    ".tsv",
    ".json",
    ".jsonl",
    ".ndjson",
    ".ipynb",
    ".md",
    ".rst",
    ".log",
    ".html",
    ".htm",
    ".xml",
    ".svg",
    ".yaml",
    ".yml",
    ".java",
    ".c",
    ".cc",
    ".cpp",
    ".cxx",
    ".h",
    ".hpp",
    ".js",
    ".jsx",
    ".ts",
    ".tsx",
    ".mjs",
    ".cjs",
    ".css",
    ".scss",
    ".sass",
    ".less",
    ".sql",
    ".r",
    ".rs",
    ".go",
    ".rb",
    ".php",
    ".swift",
    ".kt",
    ".kts",
    ".scala",
    ".lua",
    ".pl",
    ".pm",
    ".dart",
    ".ex",
    ".exs",
    ".erl",
    ".hs",
    ".clj",
    ".fs",
    ".ml",
    ".tex",
    ".bib",
    ".ini",
    ".cfg",
    ".conf",
    ".toml",
    ".env",
    ".sh",
    ".bash",
    ".zsh",
    ".fish",
    ".ps1",
    ".bat",
}

_IMAGE_MIME = {
    ".png": "image/png",
    ".jpg": "image/jpeg",
    ".jpeg": "image/jpeg",
    ".gif": "image/gif",
    ".webp": "image/webp",
}

OFFICE_SUFFIXES = {
    ".docx",
    ".docm",
    ".dotx",
    ".dotm",
    ".xlsx",
    ".xlsm",
    ".xltx",
    ".xltm",
    ".pptx",
    ".pptm",
    ".potx",
    ".potm",
    ".odt",
    ".ods",
    ".odp",
}

_DOCX_EXTS = {".docx", ".docm", ".dotx", ".dotm"}
_PPTX_EXTS = {".pptx", ".pptm", ".potx", ".potm"}
_XLSX_EXTS = {".xlsx", ".xlsm", ".xltx", ".xltm"}

_EMBEDDED_IMAGE_MIME = {
    ".png": "image/png",
    ".jpg": "image/jpeg",
    ".jpeg": "image/jpeg",
    ".gif": "image/gif",
    ".webp": "image/webp",
    ".bmp": "image/bmp",
}

MAX_TEXT_BYTES = 1_000_000
MAX_IMAGE_BYTES = 5_000_000
MAX_EMBEDDED_IMAGE_BYTES = 400_000
MAX_PDF_PAGES = 30
MAX_PPTX_SLIDES = 30
MAX_XLSX_SHEETS = 10
PDF_RESPONSE_BUDGET_BYTES = 650_000
PDF_RENDER_DPI = 110
PDF_JPEG_QUALITY = 75


def _extract_docx_markdown(filepath: str) -> str:
    import pypandoc

    return pypandoc.convert_file(
        filepath,
        "gfm",
        format="docx",
        extra_args=["--track-changes=all", "--wrap=none"],
    )


def _extract_docx_images(filepath: str, budget_remaining: int) -> tuple[list, int]:
    import base64
    import os as _os
    import zipfile

    from mcp.types import ImageContent

    blocks: list = []
    used = 0
    with zipfile.ZipFile(filepath) as z:
        names = sorted(n for n in z.namelist() if n.startswith("word/media/"))
        for name in names:
            ext = _os.path.splitext(name)[1].lower()
            mime = _EMBEDDED_IMAGE_MIME.get(ext)
            if not mime:
                continue
            blob = z.read(name)
            if len(blob) > MAX_EMBEDDED_IMAGE_BYTES:
                continue
            if used + len(blob) > budget_remaining:
                break
            used += len(blob)
            blocks.append(
                ImageContent(
                    type="image",
                    data=base64.b64encode(blob).decode(),
                    mime_type=mime,
                )
            )
    return blocks, used


def _extract_pptx_blocks(filepath: str, budget: int, pages: str | None = None) -> list:
    import base64

    from mcp.types import ImageContent, TextContent
    from pptx import Presentation

    prs = Presentation(filepath)
    all_slides = list(prs.slides)
    total = len(all_slides)

    if pages:
        indices = parse_pages(pages, total)
        truncated_by_cap = False
    else:
        indices = list(range(min(total, MAX_PPTX_SLIDES)))
        truncated_by_cap = total > MAX_PPTX_SLIDES

    blocks: list = []
    used = 0
    skipped_images = 0

    for idx in indices:
        slide = all_slides[idx]
        slide_no = idx + 1
        lines: list[str] = [f"## Slide {slide_no}"]
        images: list[tuple[bytes, str]] = []

        for shape in slide.shapes:
            if shape.has_text_frame:
                for para in shape.text_frame.paragraphs:
                    text = "".join(run.text for run in para.runs).strip()
                    if text:
                        lines.append(text)
            try:
                img = shape.image
            except AttributeError, ValueError:
                continue
            images.append((img.blob, img.content_type or "image/png"))

        if slide.has_notes_slide:
            notes = slide.notes_slide.notes_text_frame.text.strip()
            if notes:
                lines.append(f"_Notes: {notes}_")

        blocks.append(TextContent(type="text", text="\n\n".join(lines)))

        for blob, mime in images:
            if len(blob) > MAX_EMBEDDED_IMAGE_BYTES or used + len(blob) > budget:
                skipped_images += 1
                continue
            used += len(blob)
            blocks.append(
                ImageContent(
                    type="image",
                    data=base64.b64encode(blob).decode(),
                    mime_type=mime,
                )
            )

    if truncated_by_cap:
        blocks.append(
            TextContent(
                type="text",
                text=(
                    f"[TRUNCATED — showing first {MAX_PPTX_SLIDES} of "
                    f"{total} slides. Call again with "
                    f"pages='{MAX_PPTX_SLIDES + 1}-{total}' for the rest.]"
                ),
            )
        )
    if skipped_images:
        blocks.append(
            TextContent(
                type="text",
                text=(
                    f"[{skipped_images} image(s) omitted — per-image size cap "
                    f"or response budget reached.]"
                ),
            )
        )
    return blocks


def _extract_xlsx_markdown(
    filepath: str,
    pages: str | None = None,
    max_rows_per_sheet: int = 200,
) -> str:
    from openpyxl import load_workbook

    wb = load_workbook(filepath, data_only=True, read_only=True)
    try:
        all_sheets = wb.worksheets
        total = len(all_sheets)

        if pages:
            indices = parse_pages(pages, total)
            truncated_by_cap = False
        else:
            indices = list(range(min(total, MAX_XLSX_SHEETS)))
            truncated_by_cap = total > MAX_XLSX_SHEETS

        parts: list[str] = []
        for idx in indices:
            ws = all_sheets[idx]
            rows: list[tuple] = []
            for i, row in enumerate(ws.iter_rows(values_only=True)):
                if i >= max_rows_per_sheet:
                    break
                rows.append(row)

            max_cols = max((len(r) for r in rows), default=0)

            def col_empty(col: int, rows=rows) -> bool:
                return all(col >= len(r) or r[col] is None for r in rows)

            while max_cols > 0 and col_empty(max_cols - 1):
                max_cols -= 1

            parts.append(f"## Sheet: {ws.title}")
            if max_cols == 0 or not rows:
                parts.append("_(empty)_")
                continue

            def fmt(v: object) -> str:
                if v is None:
                    return ""
                return str(v).replace("|", "\\|").replace("\n", " ")

            def row_line(r: tuple, n: int = max_cols) -> str:
                cells = [fmt(r[i]) if i < len(r) else "" for i in range(n)]
                return "| " + " | ".join(cells) + " |"

            lines = [row_line(rows[0]), "| " + " | ".join(["---"] * max_cols) + " |"]
            for r in rows[1:]:
                lines.append(row_line(r))

            if ws.max_row and ws.max_row > max_rows_per_sheet:
                lines.append(
                    f"\n_[TRUNCATED — showing first {max_rows_per_sheet} of {ws.max_row} rows]_"
                )
            parts.append("\n".join(lines))

        if truncated_by_cap:
            parts.append(
                f"_[TRUNCATED — showing first {MAX_XLSX_SHEETS} of "
                f"{total} sheets. Call again with "
                f"pages='{MAX_XLSX_SHEETS + 1}-{total}' for the rest.]_"
            )
        return "\n\n".join(parts)
    finally:
        wb.close()


def parse_pages(pages: str | None, total: int) -> list[int]:
    if not pages:
        return list(range(total))
    result: set[int] = set()
    for part in pages.split(","):
        part = part.strip()
        if not part:
            continue
        if "-" in part:
            start, end = part.split("-", 1)
            for p in range(int(start), int(end) + 1):
                if 1 <= p <= total:
                    result.add(p - 1)
        else:
            p = int(part)
            if 1 <= p <= total:
                result.add(p - 1)
    return sorted(result)


def _file_to_content_inner(
    filepath: str,
    pages: str | None = None,
    display_path: str | None = None,
) -> Any:
    import base64
    import os as _os

    from mcp.types import ImageContent, TextContent

    shown_path = display_path or filepath
    filename = _os.path.basename(shown_path)
    ext = _os.path.splitext(filepath)[1].lower()
    size = _os.path.getsize(filepath)

    if ext in _TEXT_SUFFIXES:
        truncated = False
        with open(filepath, errors="replace") as f:
            text = f.read(MAX_TEXT_BYTES + 1)
        if len(text) > MAX_TEXT_BYTES:
            text = text[:MAX_TEXT_BYTES]
            truncated = True
        header = f"# {filename}\nLocal path: {shown_path}"
        if truncated:
            header += (
                f"\n[TRUNCATED — showing first {MAX_TEXT_BYTES} of {size} bytes. "
                f"Use your host's local file reader on the path above for the full file.]"
            )
        return TextContent(type="text", text=f"{header}\n\n{text}")

    if ext in _IMAGE_MIME:
        if size > MAX_IMAGE_BYTES:
            return TextContent(
                type="text",
                text=f"# {filename}\n[Image too large: {size} bytes > {MAX_IMAGE_BYTES} limit]",
            )
        with open(filepath, "rb") as f:
            data = base64.b64encode(f.read()).decode()
        return ImageContent(type="image", data=data, mime_type=_IMAGE_MIME[ext])

    if ext == ".pdf":
        import pymupdf

        doc = pymupdf.open(filepath)
        total_pages = len(doc)

        if pages:
            page_indices = parse_pages(pages, total_pages)
            header = (
                f"# {filename} ({total_pages} pages, rendering: {pages})\nLocal path: {shown_path}"
            )
        else:
            page_indices = list(range(min(total_pages, MAX_PDF_PAGES)))
            header = f"# {filename} ({total_pages} pages)\nLocal path: {shown_path}"
            if len(page_indices) < total_pages:
                header += (
                    f"\n[TRUNCATED — rendering first {len(page_indices)} of {total_pages} pages. "
                    f"Call again with pages='X-Y' to get specific pages.]"
                )

        blocks: list = [TextContent(type="text", text=header)]
        used = 0
        rendered = 0
        skipped_first: int | None = None
        for i in page_indices:
            pix = doc[i].get_pixmap(dpi=PDF_RENDER_DPI)
            img_bytes = pix.tobytes("jpeg", jpg_quality=PDF_JPEG_QUALITY)
            if used + len(img_bytes) > PDF_RESPONSE_BUDGET_BYTES and rendered > 0:
                skipped_first = i + 1
                break
            used += len(img_bytes)
            rendered += 1
            blocks.append(
                ImageContent(
                    type="image",
                    data=base64.b64encode(img_bytes).decode(),
                    mime_type="image/jpeg",
                )
            )
        doc.close()
        if skipped_first is not None:
            last_rendered = page_indices[rendered - 1] + 1
            blocks.append(
                TextContent(
                    type="text",
                    text=(
                        f"[TRUNCATED — fit {rendered} pages into the response "
                        f"size budget. Rendered through page {last_rendered}. "
                        f"Call again with pages='{skipped_first}-...' to get "
                        f"the next pages.]"
                    ),
                )
            )
        return blocks

    if ext in _DOCX_EXTS:
        md = _extract_docx_markdown(filepath)
        truncated = len(md) > MAX_TEXT_BYTES
        if truncated:
            md = md[:MAX_TEXT_BYTES]
        header = f"# {filename}\nLocal path: {shown_path}"
        if truncated:
            header += f"\n[TRUNCATED — showing first {MAX_TEXT_BYTES} chars of extracted markdown.]"
        text_block = TextContent(type="text", text=f"{header}\n\n{md}")
        image_budget = max(0, PDF_RESPONSE_BUDGET_BYTES - len(text_block.text))
        blocks = [text_block]
        image_blocks, _ = _extract_docx_images(
            filepath,
            budget_remaining=image_budget,
        )
        blocks.extend(image_blocks)
        return blocks

    if ext in _PPTX_EXTS:
        header = f"# {filename}\nLocal path: {shown_path}"
        header_block = TextContent(type="text", text=header)
        image_budget = max(0, PDF_RESPONSE_BUDGET_BYTES - len(header))
        blocks = [header_block]
        blocks.extend(_extract_pptx_blocks(filepath, budget=image_budget, pages=pages))
        return blocks

    if ext in _XLSX_EXTS:
        md = _extract_xlsx_markdown(filepath, pages=pages)
        truncated = len(md) > MAX_TEXT_BYTES
        if truncated:
            md = md[:MAX_TEXT_BYTES]
        header = f"# {filename}\nLocal path: {shown_path}"
        if truncated:
            header += f"\n[TRUNCATED — showing first {MAX_TEXT_BYTES} chars of extracted markdown.]"
        return TextContent(type="text", text=f"{header}\n\n{md}")

    return TextContent(
        type="text",
        text=(
            f"# {filename}\n"
            f"Size: {size} bytes\n"
            f"Absolute path on the user's local machine: {shown_path}\n"
            f"\n"
            f"This is a SINGLE BINARY FILE saved at the absolute path above;\n"
            f"the server did not extract it. Do NOT call `read_downloaded_file`\n"
            f"on it — that tool is only for paths returned by a previous\n"
            f"zip-listing response.\n"
            f"\n"
            f"If your host agent has a local file-read tool (Claude Code,\n"
            f"Cursor, Cline, Continue, Aider and similar all ship one), call\n"
            f"it on the absolute path above. Otherwise, ask the user to\n"
            f"upload or paste the relevant content."
        ),
    )


def file_to_content(filepath: str, pages: str | None = None) -> list:
    content = _file_to_content_inner(filepath, pages)
    return content if isinstance(content, list) else [content]
